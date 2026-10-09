#!/usr/bin/env python3
"""Recompute every QA flag and plot status in Python and compare with Excel's results.

The workbook's formulas are the product; this script is the second pair of eyes.
It reads only the input cells, applies the same rules independently, and compares
the answers with the values Excel saved. Run it after the workbook has been
opened and saved in Excel (so cached values exist):

    python qa_check.py planting-monitoring-kit-illustrative.xlsx
"""

import math
import sys
from datetime import datetime

from openpyxl import load_workbook

CHECK_ORDER = ["Unknown plot", "Visit before planting", "Count mismatch", "Alive rose without replanting",
               "Duplicate visit", "Height out of range", "No observer", "Loss without cause"]


def d(v):
    return v.date() if isinstance(v, datetime) else v


def bucket(months):
    return 1 if months < 2 else 3 if months < 4.5 else 6 if months < 7.5 else 9 if months < 10.5 else 12


def main(path):
    wb = load_workbook(path, data_only=True)
    st = wb["Settings"]
    cfg = {st.cell(row=r, column=1).value: st.cell(row=r, column=3).value for r in range(4, 12)}
    cfg["ReportDate"] = d(cfg["ReportDate"])

    pl = wb["Planting_Log"]
    plots = {}
    for r in range(5, 61):
        pid = pl.cell(row=r, column=1).value
        if pid:
            plots[pid] = {"site": pl.cell(row=r, column=2).value, "date": d(pl.cell(row=r, column=4).value),
                          "planted": pl.cell(row=r, column=7).value}

    mo = wb["Monitoring"]
    rows = []
    for r in range(5, 201):
        pid = mo.cell(row=r, column=2).value
        if not pid:
            continue
        g = lambda c: mo.cell(row=r, column=c).value
        rows.append({"r": r, "id": g(1), "plot": pid, "date": d(g(3)), "type": g(4), "obs": g(5),
                     "alive": g(6) or 0, "dead": g(7) or 0, "miss": g(8) or 0, "cause": g(9), "h": g(10),
                     "rep": g(11) or 0, "xl_flags": g(25) or "", "xl_issues": g(26)})

    for x in rows:
        p = plots.get(x["plot"])
        same = [y for y in rows if y["plot"] == x["plot"]]
        before = [y for y in same if y["date"] < x["date"]]
        x["site"] = p["site"] if p else None
        x["expected"] = p["planted"] + sum(y["rep"] for y in before) if p else None
        prev_date = max((y["date"] for y in before), default=None)
        prev_alive = sum(y["alive"] for y in same if y["date"] == prev_date) if prev_date else None
        rep_since = sum(y["rep"] for y in same if prev_date and prev_date <= y["date"] < x["date"])
        flags = []
        if not p:
            flags.append("Unknown plot")
        if p and x["date"] < p["date"]:
            flags.append("Visit before planting")
        if p and x["alive"] + x["dead"] + x["miss"] != x["expected"]:
            flags.append("Count mismatch")
        if prev_date and x["alive"] > prev_alive + rep_since:
            flags.append("Alive rose without replanting")
        if sum(1 for y in same if y["date"] == x["date"]) > 1:
            flags.append("Duplicate visit")
        if x["h"] is not None and not (cfg["HMin"] <= x["h"] <= cfg["HMax"]):
            flags.append("Height out of range")
        if not x["obs"]:
            flags.append("No observer")
        if x["dead"] + x["miss"] > 0 and not x["cause"]:
            flags.append("Loss without cause")
        x["py_flags"] = "; ".join(f for f in CHECK_ORDER if f in flags)
        x["py_issues"] = len(flags)

    checks = mismatches = 0
    for x in rows:
        for a, b in ((x["py_flags"], x["xl_flags"]), (x["py_issues"], x["xl_issues"])):
            checks += 1
            if a != b:
                mismatches += 1
                print(f"  MISMATCH {x['id']} {x['plot']}: python={a!r} excel={b!r}")

    ps = wb["Plot_Status"]
    status_rows = {}
    for r in range(5, 5 + len(plots)):
        status_rows[ps.cell(row=r, column=1).value] = {c: ps.cell(row=r, column=c).value for c in range(1, 20)}
    for pid, p in plots.items():
        clean = [x for x in rows if x["plot"] == pid and x["py_issues"] == 0]
        last = max(x["date"] for x in clean)
        months = (last - p["date"]).days / 30.44
        alive = sum(x["alive"] for x in clean if x["date"] == last)
        rep_before = sum(x["rep"] for x in rows if x["plot"] == pid and x["date"] < last)
        rep_at = sum(x["rep"] for x in rows if x["plot"] == pid and x["date"] == last)
        stocking = alive / p["planted"]
        if rep_at > 0:
            status = "Replanted, recheck"
        elif bucket(months) >= cfg["TriggerMonth"] and stocking < cfg["Trigger"]:
            status = "Replant"
        elif stocking < cfg["Trigger"] + cfg["WatchBand"]:
            status = "Watch"
        else:
            status = "On track"
        to_replant = math.ceil(p["planted"] * cfg["Target"] - alive - 1e-9) if status == "Replant" else 0
        days = (cfg["ReportDate"] - last).days
        open_qa = sum(1 for x in rows if x["plot"] == pid and x["py_issues"] > 0)
        if status == "Replant":
            action = f"Replant {to_replant} seedlings with the partner group" + ("; census overdue" if days > cfg["OverdueDays"] else "")
        elif days > cfg["OverdueDays"]:
            action = "Census overdue"
        elif status == "Replanted, recheck":
            action = "Recount at the next census"
        elif status == "Watch":
            action = "Check the cause at the next visit"
        else:
            action = "Routine census"
        if open_qa:
            action += f"; fix {open_qa} flagged row" + ("s" if open_qa > 1 else "")
        cause = next(x["cause"] for x in clean if x["date"] == last) or ""
        xl = status_rows[pid]
        expect = {5: last, 6: bucket(months), 7: alive, 8: rep_before, 9: round(stocking, 6),
                  10: round(alive / (p["planted"] + rep_before), 6), 11: rep_at, 12: status, 13: to_replant,
                  14: days, 15: open_qa, 16: action, 17: p["planted"] + rep_before - alive, 18: cause}
        for col, val in expect.items():
            got = xl[col]
            got = d(got) if isinstance(got, datetime) else got
            if isinstance(val, float):
                got = round(got, 6)
            checks += 1
            if got != val:
                mismatches += 1
                print(f"  MISMATCH Plot_Status {pid} col {col}: python={val!r} excel={got!r}")

    flagged = sum(1 for x in rows if x["py_issues"] > 0)
    print(f"rows {len(rows)}, plots {len(plots)}, flagged rows {flagged}")
    print(f"checks {checks}, mismatches {mismatches}")
    return mismatches


if __name__ == "__main__":
    sys.exit(1 if main(sys.argv[1]) else 0)
