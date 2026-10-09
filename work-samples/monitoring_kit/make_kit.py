#!/usr/bin/env python3
"""Build the Planting and Monitoring Data Kit workbook (illustrative data).

ILLUSTRATIVE DATA. Every site, plot, count and date below is invented to show
how the sheets work. Nothing here describes a real CarbonEthics site.

Sheets
  Read_Me         how to use the kit, data rules, the eight QA checks
  Settings        report date, replanting trigger, plausible ranges, lists
  Planting_Log    one row per plot planted (inputs)
  Monitoring      one row per plot visit (inputs) + formula columns and QA flags
  Plot_Status     one row per plot: last clean census, stocking, status, next action
  Site_Summary    per-site totals, stocking trend, losses by main cause
  Dashboard       key figures and three charts
  Field_Day_SOP   planting-day procedure, owner and record for every step
  Risk_Register   12 risks scored likelihood x severity, with controls

Run:  python make_kit.py <output.xlsx>
"""

import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.data_source import StrRef
from openpyxl.chart.series import SeriesLabel
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

ACCENT = "0F5C6E"
INPUT_FILL = PatternFill("solid", start_color="FFF8E1", end_color="FFF8E1")
CALC_FILL = PatternFill("solid", start_color="F2F4F7", end_color="F2F4F7")
HEAD_FILL = PatternFill("solid", start_color=ACCENT, end_color=ACCENT)
BAD = PatternFill("solid", start_color="FDE2E1", end_color="FDE2E1")
WARN = PatternFill("solid", start_color="FFF1CC", end_color="FFF1CC")
GOOD = PatternFill("solid", start_color="E3F4EA", end_color="E3F4EA")
INFO = PatternFill("solid", start_color="E8F0FB", end_color="E8F0FB")
THIN = Side(style="thin", color="D0D5DD")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HFONT = Font(name="Calibri", bold=True, color="FFFFFF", size=10)
BFONT = Font(name="Calibri", size=10)
TITLE = Font(name="Calibri", bold=True, size=15, color="101828")
SUB = Font(name="Calibri", italic=True, size=10, color="475467")
S1, S2, S3, GREY = "2A78D6", "EB6834", "1BAF7A", "52514E"

LAST = 200          # formulas reference rows 5..LAST
FORMULA_ROWS = 80   # Monitoring rows that carry formulas (data + spare rows)
DATE_FMT = "dd mmm yyyy"

SITES = ["Site A - seaward fringe", "Site B - pond dyke", "Site C - sheltered lagoon"]
SPECIES = ["Rhizophora mucronata", "Rhizophora apiculata", "Avicennia marina", "Sonneratia alba"]
CAUSES = ["Wave/scour", "Barnacles", "Sediment burial", "Crab damage", "Grazing",
          "Debris/plastic", "Human disturbance", "Unknown"]
OBSERVERS = ["FS-1 (field staff)", "FS-2 (field staff)", "PL-1 (partner lead)",
             "PL-2 (partner lead)", "PL-3 (partner lead)"]
VISIT_TYPES = ["Census", "Replanting", "Follow-up"]

PLOTS = [
    # plot, site, zone, planted on, species, source, seedlings, spacing m, partner, waypoint, photo point
    ("A01", 0, "Open fringe", date(2025, 10, 18), 0, "Partner nursery", 500, 1.0, "Partner group 1"),
    ("A02", 0, "Open fringe", date(2025, 10, 18), 0, "Partner nursery", 500, 1.0, "Partner group 1"),
    ("A03", 0, "Behind wave breaker", date(2025, 10, 18), 0, "Partner nursery", 500, 1.0, "Partner group 1"),
    ("A04", 0, "Behind wave breaker", date(2025, 10, 19), 0, "Partner nursery", 450, 1.0, "Partner group 1"),
    ("B01", 1, "Dyke crest", date(2026, 1, 24), 0, "Partner nursery", 600, 1.0, "Partner group 2"),
    ("B02", 1, "Dyke crest", date(2026, 1, 24), 0, "Partner nursery", 600, 1.0, "Partner group 2"),
    ("B03", 1, "Pond floor", date(2026, 1, 24), 0, "Partner nursery", 550, 1.0, "Partner group 2"),
    ("B04", 1, "Pond floor", date(2026, 1, 24), 0, "Partner nursery", 600, 1.0, "Partner group 2"),
    ("C01", 2, "Lagoon edge", date(2026, 3, 14), 2, "Wild propagules, partner nursery", 400, 1.5, "Partner group 3"),
    ("C02", 2, "Lagoon edge", date(2026, 3, 14), 2, "Wild propagules, partner nursery", 400, 1.5, "Partner group 3"),
    ("C03", 2, "Inner flat", date(2026, 3, 14), 2, "Wild propagules, partner nursery", 400, 1.5, "Partner group 3"),
    ("C04", 2, "Inner flat", date(2026, 3, 14), 2, "Wild propagules, partner nursery", 400, 1.5, "Partner group 3"),
]

# plot, visit date, type, observer index, alive, dead, missing, cause index or None, height cm, replanted, note
VISITS = [
    ("A01", date(2025, 11, 18), "Census", 0, 455, 30, 15, 0, 48, 0, ""),
    ("A01", date(2026, 1, 17), "Census", 0, 360, 90, 50, 0, 55, 0, ""),
    ("A01", date(2026, 4, 18), "Census", 0, 290, 130, 80, 0, 64, 0, "Scour on the seaward edge"),
    ("A01", date(2026, 5, 16), "Replanting", 2, 286, 132, 82, 0, 66, 115, "Replanted with stakes"),
    ("A02", date(2025, 11, 18), "Census", 0, 440, 40, 20, 0, 47, 0, ""),
    ("A02", date(2026, 1, 17), "Census", 0, 330, 110, 60, 1, 54, 0, ""),
    ("A02", date(2026, 4, 18), "Census", 0, 245, 150, 105, 1, 62, 0, "Heavy barnacle load"),
    ("A02", date(2026, 5, 16), "Replanting", 2, 243, 152, 105, 1, 63, 157, "Stems cleaned before replanting"),
    ("A02", date(2026, 7, 18), "Follow-up", 1, 372, 170, 115, 1, 70, 0, "Replanted seedlings mostly alive"),
    ("A03", date(2025, 11, 18), "Census", 1, 470, 20, 10, 2, 49, 0, ""),
    ("A03", date(2026, 1, 17), "Census", 1, 410, 60, 50, 2, 56, 0, "Tally sheet re-added on the boat"),
    ("A03", date(2026, 4, 18), "Census", 1, 365, 95, 40, 1, 66, 0, ""),
    ("A04", date(2025, 11, 19), "Census", 1, 425, 15, 10, 0, 48, 0, ""),
    ("A04", date(2026, 1, 17), "Census", 1, 380, 45, 25, 0, 55, 0, ""),
    ("A04", date(2026, 4, 18), "Census", 1, 300, 95, 55, 1, 65, 0, ""),
    ("B01", date(2026, 2, 24), "Census", 3, 575, 15, 10, None, 46, 0, ""),
    ("B01", date(2026, 4, 25), "Census", 3, 540, 35, 25, 3, 52, 0, ""),
    ("B01", date(2026, 7, 25), "Census", 0, 510, 55, 35, 3, 63, 0, ""),
    ("B02", date(2026, 2, 24), "Census", 3, 580, 12, 8, 3, 47, 0, ""),
    ("B02", date(2026, 4, 25), "Census", 3, 545, 35, 20, 3, 53, 0, ""),
    ("B02", date(2026, 7, 25), "Census", 0, 562, 23, 15, 3, 62, 0, ""),
    ("B03", date(2026, 2, 24), "Census", 3, 520, 20, 10, 2, 45, 0, ""),
    ("B3", date(2026, 4, 25), "Census", 3, 455, 60, 35, 2, 51, 0, ""),
    ("B03", date(2026, 7, 25), "Census", 0, 365, 115, 70, 2, 58, 0, "Silt over the lower rows"),
    ("B04", date(2026, 2, 24), "Census", 3, 585, 10, 5, 3, 46, 0, ""),
    ("B04", date(2026, 4, 25), "Census", 3, 560, 25, 15, 3, 850, 0, ""),
    ("B04", date(2026, 7, 25), "Census", 0, 535, 40, 25, 3, 63, 0, ""),
    ("C01", date(2025, 4, 14), "Census", 4, 388, 8, 4, 5, 30, 0, ""),
    ("C01", date(2026, 6, 13), "Census", 4, 365, 22, 13, 5, 38, 0, ""),
    ("C01", date(2026, 9, 12), "Census", 1, 340, 38, 22, 5, 50, 0, "Plastic caught on seedlings"),
    ("C02", date(2026, 4, 14), "Census", 4, 380, 12, 8, 4, 31, 0, "Goats seen near the edge"),
    ("C02", date(2026, 6, 13), "Census", 4, 330, 45, 25, 4, 37, 0, ""),
    ("C02", date(2026, 9, 12), "Census", 1, 290, 70, 40, 4, 49, 0, ""),
    ("C03", date(2026, 4, 14), "Census", 4, 392, 5, 3, 5, 32, 0, ""),
    ("C03", date(2026, 6, 13), "Census", 4, 372, 18, 10, 5, 39, 0, ""),
    ("C03", date(2026, 6, 13), "Census", 4, 372, 18, 10, 5, 39, 0, "Entered from photo of the sheet"),
    ("C03", date(2026, 9, 12), "Census", 1, 355, 30, 15, 5, 51, 0, ""),
    ("C04", date(2026, 4, 14), "Census", 4, 390, 6, 4, 3, 31, 0, ""),
    ("C04", date(2026, 6, 13), "Census", 4, 375, 15, 10, 3, 38, 0, ""),
    ("C04", date(2026, 9, 12), "Census", None, 362, 25, 13, 3, 50, 0, ""),
]

SOP = [
    ("T-14 to T-7: plan", "Confirm the plots to be planted, their boundaries and waypoint IDs", "Every plot has an ID in Planting_Log before the day", "Ops intern", "Planting_Log rows"),
    ("T-14 to T-7: plan", "Choose the date and start time from the tide table", "Planting on a falling-to-low tide; crew off the flat before the flood returns", "Field lead", "Tide window in the day plan"),
    ("T-14 to T-7: plan", "Check that each zone is plantable", "Elevation and flooding compared with nearby natural mangroves; species matched to the zone", "Science team", "Site note with photos"),
    ("T-14 to T-7: plan", "Agree roles, numbers and compensation with the partner group in writing", "Who plants, who maintains, how much and when, in Bahasa Indonesia", "Ops lead + partner lead", "Signed agreement or written confirmation on file"),
    ("T-14 to T-7: plan", "Order seedlings by zone", "Planned count plus 10% spare; minimum height and leaf count stated in the order", "Ops intern", "Seedling order"),
    ("T-14 to T-7: plan", "Permits and notifications", "Village office or area authority informed as required; client briefing pack for corporate plantings", "Ops lead", "Letter or message on file"),
    ("T-14 to T-7: plan", "Participant list and safety kit", "Emergency contacts; boots, gloves, first aid, water; life jackets for any boat transfer", "Ops intern", "Participant list, kit checklist"),
    ("T-1: check", "Seedling quality check at the nursery", "Count, height, healthy leaves, intact root ball, no pests; reject below standard", "Science team + partner", "QC sheet: accepted and rejected"),
    ("T-1: check", "Weather and sea check", "Marine forecast checked; go or no-go sent by 06:00 on the day; backup date agreed", "Field lead", "Go/no-go message"),
    ("T-1: check", "Prepare data tools", "Forms loaded, plot tags with batch codes printed, GPS and phones charged", "Ops intern", "Ready checklist"),
    ("Day: plant", "Safety and method briefing", "Tide time, mud safety, planting depth, staking and tying, nothing plastic left behind", "Field lead", "Briefing done before planting starts"),
    ("Day: plant", "Assign plots to teams, one recorder per plot", "Each recorder tallies seedlings planted in their plot", "Ops intern", "Tally per plot"),
    ("Day: plant", "Take photo-point pictures", "Same position and direction on every visit, plot ID in the file name", "Ops intern", "Photo log"),
    ("Day: plant", "Collect polybags and ties", "Bags out equals bags back", "Everyone", "Bag count"),
    ("Close-out (same day)", "Three-way seedling count", "Delivered = planted + spare returned + rejected; signed by field lead and partner lead", "Ops intern", "Signed tally"),
    ("Close-out (same day)", "Enter the data and run the checks", "Rows entered the same day, paper sheets photographed, QA flags at zero", "Ops intern", "Planting_Log and Monitoring rows"),
    ("Follow-up", "Census at months 1, 3, 6 and 12", "Replanting by the agreed terms when stocking falls below the trigger", "Science team + partner", "Monitoring rows, Plot_Status actions"),
    ("Follow-up", "Pay the partner group on the agreed date and share results back", "Payment logged; results explained to the group in a short session", "Ops lead + finance", "Payment log, session note"),
]

RISKS = [
    ("R01", "Seedlings planted in the wrong zone (too low, flooded too long)", "Site chosen by convenience, not by flooding", 2, 3, "Compare elevation and flooding with nearby natural mangroves; pilot plot before scaling", "Plant only inside the marked zone", "Science team", "Month-1 losses above 20% in the lowest plots"),
    ("R02", "Tide window missed", "Ceremony or boat delay", 2, 3, "Schedule from the tide table; ceremony after planting or kept short", "Hard stop time announced at the briefing", "Field lead", "Start delayed by more than 30 minutes"),
    ("R03", "Wave exposure and scour on open fringe", "Exposed plots, no shelter", 2, 3, "Plant sheltered zones first, behind wave breakers where they exist; stake seedlings", "Check stakes and ties before leaving", "Science team", "Missing counts rising at month 1"),
    ("R04", "Barnacle load and debris on stems", "Low maintenance between visits", 2, 2, "Monthly cleaning duty written into the partner agreement", "Show cleaning method at the briefing", "Partner lead", "Barnacles named as main cause at a census"),
    ("R05", "Poor seedling quality", "Too young, damaged or diseased stock", 2, 2, "Nursery check on T-1 with written reject criteria; 10% spare", "Recorder sets aside rejected seedlings", "Science team", "Rejection rate above 10%"),
    ("R06", "Counts do not reconcile (delivered, planted, reported)", "Tally errors, late data entry", 2, 3, "One recorder per plot; tally sheets printed", "Three-way count signed; same-day entry with QA checks", "Ops intern", "QA flags on entry"),
    ("R07", "Plot cannot be found again", "No marker, GPS drift, photo point not fixed", 2, 2, "Fixed markers, waypoint ID and photo point for every plot", "Photo of each marker", "Ops intern", "Census note: plot not located"),
    ("R08", "Partner group unclear on roles or payment", "Terms agreed verbally only", 2, 3, "Written agreement in Bahasa Indonesia; named contact; payment date", "Repeat the terms at the briefing", "Ops lead", "Maintenance visits missed"),
    ("R09", "Grazing, trampling or human disturbance", "Livestock or foot traffic near the site", 1, 2, "Agree signs or simple fencing with the village where livestock roam", "Mark planted rows clearly", "Partner lead", "Grazing named as main cause"),
    ("R10", "Injury: cuts, slips, heat, boat transfer", "Mud, shells, sun, small boats", 1, 3, "PPE list, first-aid kit, life jackets, buddy system", "Safety briefing; water breaks", "Field lead", "Any near-miss reported"),
    ("R11", "Bad weather or high waves", "Season, storms", 2, 2, "Marine forecast check; go or no-go by 06:00; backup date", "Stop work if waves or lightning", "Field lead", "Forecast warning issued"),
    ("R12", "Plastic left on site", "Polybags and ties dropped", 2, 1, "Bag count planned", "Bags out equals bags back", "Everyone", "Bag count short"),
]


def head(ws, row, labels, widths=None, height=30):
    for i, text in enumerate(labels, start=1):
        c = ws.cell(row=row, column=i, value=text)
        c.font, c.fill, c.border = HFONT, HEAD_FILL, BOX
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[row].height = height
    if widths:
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w


def title(ws, text, sub):
    ws["A1"] = text
    ws["A1"].font = TITLE
    ws["A2"] = sub
    ws["A2"].font = SUB


def style_range(ws, r1, r2, c1, c2, fill=None, fmt=None, wrap=False):
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = BOX
            cell.font = BFONT
            if fill:
                cell.fill = fill
            if fmt:
                cell.number_format = fmt
            cell.alignment = Alignment(vertical="top", wrap_text=wrap)


def page(ws, landscape=True):
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5


def add_list_validation(ws, ref, cells):
    dv = DataValidation(type="list", formula1=ref, allow_blank=True, showErrorMessage=True,
                        errorTitle="Not on the list", error="Pick a value from the list in Settings.")
    ws.add_data_validation(dv)
    dv.add(cells)


def build(out_path):
    wb = Workbook()

    # ------------------------------------------------------------------ Read_Me
    rm = wb.active
    rm.title = "Read_Me"
    title(rm, "Planting and Monitoring Data Kit (illustrative data)",
          "Built by Les' Aullian Achmad Argito for the CarbonEthics Operations & Science Intern application, October 2026.")
    lines = [
        ("Read this first", "Every site, plot, count and date in this file is invented to show how the sheets work. Nothing here describes a real CarbonEthics site, partner or client."),
        ("What it is for", "Keeping planting and survival data clean enough to report: one row per plot planted, one row per plot visit, automatic checks, and a status line per plot that tells the team what to do next."),
        ("How to use it", "1) Set the report date and thresholds in Settings. 2) Add plots to Planting_Log. 3) Add one Monitoring row per plot visit (yellow cells only; grey cells are formulas). 4) Read the QA_Flags column and correct flagged rows at the source. 5) Read Plot_Status and Dashboard for actions. 6) Copy grey formula rows down when you add rows."),
        ("Data rule 1", "Counts are cumulative since planting: Alive + Dead + Missing must equal seedlings planted plus seedlings replanted on earlier visits."),
        ("Data rule 2", "Replanting is recorded on the visit where it happens, after that visit's count; it enters the expected total from the next visit."),
        ("Data rule 3", "Plot_Status uses only rows that pass every check. A flagged row is corrected, never deleted, and the plot shows its last clean census until then."),
        ("Data rule 4", "Stocking = living trees / seedlings in the original design (drives the replanting trigger). Survival = living trees / all seedlings ever planted on the plot (shows what replanting cost)."),
        ("QA check 1: Unknown plot", "Plot ID not found in Planting_Log (typos such as B3 for B03)."),
        ("QA check 2: Visit before planting", "Visit date earlier than the planting date (often a mistyped year)."),
        ("QA check 3: Count mismatch", "Alive + Dead + Missing does not equal the expected total."),
        ("QA check 4: Alive rose without replanting", "More living trees than at the previous visit plus anything replanted since."),
        ("QA check 5: Duplicate visit", "Same plot and date entered twice."),
        ("QA check 6: Height out of range", "Mean height outside the plausible range in Settings (a slipped decimal point)."),
        ("QA check 7: No observer", "Nobody named as responsible for the count."),
        ("QA check 8: Loss without cause", "Dead or missing seedlings recorded with no main cause."),
        ("Dropdowns", "Lists in Settings stop most typos at entry; the checks catch what arrives by import, copy-paste or photos of paper sheets."),
        ("Google Sheets", "Uses COUNTIFS, SUMIFS, AVERAGEIFS, SUMPRODUCT and INDEX/MATCH only, which work in Google Sheets and Excel."),
    ]
    rm.column_dimensions["A"].width = 32
    rm.column_dimensions["B"].width = 110
    for i, (k, v) in enumerate(lines, start=4):
        rm.cell(row=i, column=1, value=k).font = Font(name="Calibri", bold=True, size=10, color=ACCENT)
        c = rm.cell(row=i, column=2, value=v)
        c.font = BFONT
        c.alignment = Alignment(wrap_text=True, vertical="top")
        rm.cell(row=i, column=1).alignment = Alignment(vertical="top", wrap_text=True)
    page(rm, landscape=False)

    # ------------------------------------------------------------------ Settings
    st = wb.create_sheet("Settings")
    title(st, "Settings", "Change these values; every formula reads them by name.")
    settings = [
        ("ReportDate", "Report date", date(2026, 9, 30), DATE_FMT),
        ("Trigger", "Replant when stocking is below", 0.70, "0%"),
        ("TriggerMonth", "Apply the trigger from census month", 6, "0"),
        ("WatchBand", "Watch band above the trigger", 0.10, "0%"),
        ("Target", "Replant up to this stocking", 0.80, "0%"),
        ("HMin", "Plausible mean height, minimum (cm)", 10, "0"),
        ("HMax", "Plausible mean height, maximum (cm)", 250, "0"),
        ("OverdueDays", "Census overdue after (days)", 100, "0"),
    ]
    head(st, 3, ["Name", "Setting", "Value"], [16, 38, 16])
    for i, (name, label, value, fmt) in enumerate(settings, start=4):
        st.cell(row=i, column=1, value=name)
        st.cell(row=i, column=2, value=label)
        st.cell(row=i, column=3, value=value).number_format = fmt
        wb.defined_names[name] = DefinedName(name, attr_text=f"Settings!$C${i}")
    style_range(st, 4, 3 + len(settings), 1, 2, CALC_FILL)
    style_range(st, 4, 3 + len(settings), 3, 3, INPUT_FILL)
    for i, (_, _, _, fmt) in enumerate(settings, start=4):
        st.cell(row=i, column=3).number_format = fmt
    lists = [("Sites", SITES, "E"), ("Species", SPECIES, "F"), ("Causes", CAUSES, "G"),
             ("Observers", OBSERVERS, "H"), ("Visit types", VISIT_TYPES, "I")]
    for name, values, col in lists:
        c = st[f"{col}3"]
        c.value, c.font, c.fill, c.border = name, HFONT, HEAD_FILL, BOX
        st.column_dimensions[col].width = 26
        for j, v in enumerate(values, start=4):
            cell = st[f"{col}{j}"]
            cell.value, cell.border, cell.font, cell.fill = v, BOX, BFONT, INPUT_FILL
    page(st)

    # ------------------------------------------------------------------ Planting_Log
    pl = wb.create_sheet("Planting_Log")
    title(pl, "Planting log", "One row per plot planted. Yellow = input. Illustrative data.")
    head(pl, 4, ["Plot ID", "Site", "Zone", "Planted on", "Species", "Seedling source", "Seedlings planted",
                 "Spacing (m)", "Partner group", "GPS waypoint ID", "Photo point ID", "Notes"],
         [9, 26, 20, 13, 22, 28, 11, 9, 16, 13, 13, 24])
    for r, p in enumerate(PLOTS, start=5):
        plot, site, zone, d, sp, src, n, spacing, partner = p
        values = [plot, SITES[site], zone, d, SPECIES[sp], src, n, spacing, partner, f"WP-{plot}", f"PP-{plot}", ""]
        for c, v in enumerate(values, start=1):
            pl.cell(row=r, column=c, value=v)
    last_pl = 4 + len(PLOTS)
    style_range(pl, 5, last_pl + 6, 1, 12, INPUT_FILL)
    for r in range(5, last_pl + 7):
        pl.cell(row=r, column=4).number_format = DATE_FMT
    add_list_validation(pl, "=Settings!$E$4:$E$6", f"B5:B{last_pl + 6}")
    add_list_validation(pl, "=Settings!$F$4:$F$7", f"E5:E{last_pl + 6}")
    pl.freeze_panes = "B5"
    page(pl)

    # ------------------------------------------------------------------ Monitoring
    mo = wb.create_sheet("Monitoring")
    title(mo, "Monitoring census", "One row per plot visit. Yellow = input, grey = formula. Illustrative data with eight deliberate entry errors for the checks to catch.")
    mo_head = ["Visit ID", "Plot ID", "Visit date", "Visit type", "Observer", "Alive", "Dead", "Missing",
               "Main cause of loss", "Mean height (cm)", "Replanted this visit", "Notes",
               "Site", "Planted on", "Months since planting", "Census month", "Seedlings planted",
               "Replanted before this visit", "Expected total", "Survival (all planted)", "Stocking (vs design)",
               "Previous visit", "Alive at previous visit", "Replanted since previous", "QA flags", "QA issues"]
    head(mo, 4, mo_head, [8, 8, 12, 11, 18, 7, 7, 8, 16, 9, 10, 26,
                          24, 12, 9, 8, 9, 10, 9, 10, 10, 12, 9, 10, 40, 7], height=44)
    for r, v in enumerate(VISITS, start=5):
        plot, d, vtype, obs, alive, dead, miss, cause, h, rep, note = v
        row = [f"V{r - 4:03d}", plot, d, vtype, OBSERVERS[obs] if obs is not None else None,
               alive, dead, miss, CAUSES[cause] if cause is not None else None, h, rep, note]
        for c, val in enumerate(row, start=1):
            mo.cell(row=r, column=c, value=val)
    rng = lambda col: f"${col}$5:${col}${LAST}"
    for r in range(5, 5 + FORMULA_ROWS):
        f = {
            "M": f'=IF($B{r}="","",IFERROR(INDEX(Planting_Log!$B$5:$B$60,MATCH($B{r},Planting_Log!$A$5:$A$60,0)),""))',
            "N": f'=IF($B{r}="","",IFERROR(INDEX(Planting_Log!$D$5:$D$60,MATCH($B{r},Planting_Log!$A$5:$A$60,0)),""))',
            "O": f'=IF(OR(N{r}="",C{r}=""),"",ROUND((C{r}-N{r})/30.44,1))',
            "P": f'=IF(O{r}="","",IF(O{r}<2,1,IF(O{r}<4.5,3,IF(O{r}<7.5,6,IF(O{r}<10.5,9,12)))))',
            "Q": f'=IF($B{r}="","",IFERROR(INDEX(Planting_Log!$G$5:$G$60,MATCH($B{r},Planting_Log!$A$5:$A$60,0)),""))',
            "R": f'=IF($B{r}="","",SUMIFS({rng("K")},{rng("B")},$B{r},{rng("C")},"<"&$C{r}))',
            "S": f'=IF(Q{r}="","",Q{r}+R{r})',
            "T": f'=IF(OR(S{r}="",S{r}=0),"",F{r}/S{r})',
            "U": f'=IF(OR(Q{r}="",Q{r}=0),"",F{r}/Q{r})',
            "V": f'=IF($B{r}="","",IF(COUNTIFS({rng("B")},$B{r},{rng("C")},"<"&$C{r})=0,"",SUMPRODUCT(MAX(({rng("B")}=$B{r})*({rng("C")}<$C{r})*{rng("C")}))))',
            "W": f'=IF(V{r}="","",SUMIFS({rng("F")},{rng("B")},$B{r},{rng("C")},V{r}))',
            "X": f'=IF($B{r}="","",IF(V{r}="",0,SUMIFS({rng("K")},{rng("B")},$B{r},{rng("C")},">="&V{r},{rng("C")},"<"&$C{r})))',
            # nested IFs, not AND(): Excel's AND evaluates every argument, so a blank
            # previous-visit cell would turn the whole check into #VALUE!
            "Y": (f'=IF($B{r}="","",MID('
                  f'IF(M{r}="","; Unknown plot","")'
                  f'&IF(N{r}="","",IF(C{r}<N{r},"; Visit before planting",""))'
                  f'&IF(S{r}="","",IF(F{r}+G{r}+H{r}<>S{r},"; Count mismatch",""))'
                  f'&IF(W{r}="","",IF(F{r}>W{r}+X{r},"; Alive rose without replanting",""))'
                  f'&IF(COUNTIFS({rng("B")},$B{r},{rng("C")},$C{r})>1,"; Duplicate visit","")'
                  f'&IF(J{r}="","",IF(OR(J{r}<HMin,J{r}>HMax),"; Height out of range",""))'
                  f'&IF(E{r}="","; No observer","")'
                  f'&IF(G{r}+H{r}>0,IF(I{r}="","; Loss without cause",""),"")'
                  f',3,300))'),
            "Z": f'=IF($B{r}="","",IF(Y{r}="",0,1+LEN(Y{r})-LEN(SUBSTITUTE(Y{r},";",""))))',
        }
        for col, formula in f.items():
            mo[f"{col}{r}"] = formula
    end = 4 + FORMULA_ROWS
    style_range(mo, 5, end, 1, 12, INPUT_FILL)
    style_range(mo, 5, end, 13, 26, CALC_FILL)
    for r in range(5, end + 1):
        mo[f"C{r}"].number_format = DATE_FMT
        mo[f"N{r}"].number_format = DATE_FMT
        mo[f"V{r}"].number_format = DATE_FMT
        mo[f"T{r}"].number_format = "0%"
        mo[f"U{r}"].number_format = "0%"
        mo[f"Y{r}"].alignment = Alignment(wrap_text=True, vertical="top")
    mo.conditional_formatting.add(f"A5:Z{end}", FormulaRule(formula=[f"AND($Z5<>\"\",$Z5>0)"], fill=BAD))
    add_list_validation(mo, "=Planting_Log!$A$5:$A$60", f"B5:B{end}")
    add_list_validation(mo, "=Settings!$I$4:$I$6", f"D5:D{end}")
    add_list_validation(mo, "=Settings!$H$4:$H$8", f"E5:E{end}")
    add_list_validation(mo, "=Settings!$G$4:$G$11", f"I5:I{end}")
    mo.freeze_panes = "C5"
    page(mo)

    # ------------------------------------------------------------------ Plot_Status
    ps = wb.create_sheet("Plot_Status")
    title(ps, "Plot status", "One row per plot, from the last census that passed every check. Read Next action first.")
    ps_head = ["Plot ID", "Site", "Seedlings planted", "Planted on", "Last clean visit", "Census month at last visit",
               "Alive at last visit", "Replanted before last visit", "Stocking (vs design)", "Survival (all planted)",
               "Replanted at last visit", "Status", "Seedlings to replant", "Days since last visit",
               "Open QA rows", "Next action", "Seedlings lost", "Main cause at last visit", "Trigger"]
    head(ps, 4, ps_head, [8, 24, 10, 12, 12, 9, 9, 11, 10, 10, 10, 17, 10, 9, 8, 46, 9, 16, 8], height=44)
    M = "Monitoring!"
    mr = lambda col: f"{M}${col}$5:${col}${LAST}"
    for i in range(len(PLOTS)):
        r = 5 + i
        pr = 5 + i
        f = {
            "A": f'=IF(Planting_Log!A{pr}="","",Planting_Log!A{pr})',
            "B": f'=IF(A{r}="","",Planting_Log!B{pr})',
            "C": f'=IF(A{r}="","",Planting_Log!G{pr})',
            "D": f'=IF(A{r}="","",Planting_Log!D{pr})',
            "E": (f'=IF(A{r}="","",IF(SUMPRODUCT(({mr("B")}=A{r})*({mr("Z")}=0)*({mr("C")}<>""))=0,"",'
                  f'SUMPRODUCT(MAX(({mr("B")}=A{r})*({mr("Z")}=0)*{mr("C")}))))'),
            "F": (f'=IF(E{r}="","",IF((E{r}-D{r})/30.44<2,1,IF((E{r}-D{r})/30.44<4.5,3,'
                  f'IF((E{r}-D{r})/30.44<7.5,6,IF((E{r}-D{r})/30.44<10.5,9,12)))))'),
            "G": f'=IF(E{r}="","",SUMIFS({mr("F")},{mr("B")},A{r},{mr("C")},E{r},{mr("Z")},0))',
            "H": f'=IF(E{r}="","",SUMIFS({mr("K")},{mr("B")},A{r},{mr("C")},"<"&E{r}))',
            "I": f'=IF(E{r}="","",G{r}/C{r})',
            "J": f'=IF(E{r}="","",G{r}/(C{r}+H{r}))',
            "K": f'=IF(E{r}="","",SUMIFS({mr("K")},{mr("B")},A{r},{mr("C")},E{r}))',
            "L": (f'=IF(A{r}="","",IF(E{r}="","No clean census",IF(K{r}>0,"Replanted, recheck",'
                  f'IF(AND(F{r}>=TriggerMonth,I{r}<Trigger),"Replant",IF(I{r}<Trigger+WatchBand,"Watch","On track")))))'),
            "M": f'=IF(L{r}="Replant",ROUNDUP(C{r}*Target-G{r},0),0)',
            "N": f'=IF(E{r}="","",ReportDate-E{r})',
            "O": f'=IF(A{r}="","",COUNTIFS({mr("B")},A{r},{mr("Z")},">0"))',
            "P": (f'=IF(A{r}="","",IF(L{r}="No clean census","Correct flagged rows, then recount",'
                  f'IF(L{r}="Replant","Replant "&M{r}&" seedlings with the partner group"&IF(N{r}>OverdueDays,"; census overdue",""),'
                  f'IF(N{r}>OverdueDays,"Census overdue",IF(L{r}="Replanted, recheck","Recount at the next census",'
                  f'IF(L{r}="Watch","Check the cause at the next visit","Routine census")))))'
                  f'&IF(O{r}>0,"; fix "&O{r}&" flagged row"&IF(O{r}>1,"s",""),""))'),
            "Q": f'=IF(E{r}="","",C{r}+H{r}-G{r})',
            "R": (f'=IF(E{r}="","",INDEX({mr("I")},SUMPRODUCT(({mr("B")}=A{r})*({mr("C")}=E{r})*({mr("Z")}=0)'
                  f'*(ROW({mr("B")})-4)))&"")'),
            "S": f'=IF(A{r}="","",Trigger)',
        }
        for col, formula in f.items():
            ps[f"{col}{r}"] = formula
    pend = 4 + len(PLOTS)
    style_range(ps, 5, pend, 1, 19, CALC_FILL)
    for r in range(5, pend + 1):
        for col in ("D", "E"):
            ps[f"{col}{r}"].number_format = DATE_FMT
        for col in ("I", "J", "S"):
            ps[f"{col}{r}"].number_format = "0%"
        ps[f"P{r}"].alignment = Alignment(wrap_text=True, vertical="top")
    for text, fill in (("Replant", BAD), ("Watch", WARN), ("On track", GOOD), ("Replanted, recheck", INFO),
                       ("No clean census", BAD)):
        ps.conditional_formatting.add(f"L5:L{pend}", FormulaRule(formula=[f'$L5="{text}"'], fill=fill))
    ps.freeze_panes = "B5"
    page(ps)

    # ------------------------------------------------------------------ Site_Summary
    ss = wb.create_sheet("Site_Summary")
    title(ss, "Site summary", "Totals per site, stocking by census month, and losses by main cause.")
    head(ss, 4, ["Site", "Plots", "Seedlings planted", "Replanted to date", "Alive at last clean visit",
                 "Stocking (vs design)", "Plots to replant", "Seedlings to order", "Plots with census overdue",
                 "Flagged rows open"], [26, 7, 11, 11, 12, 11, 10, 11, 11, 10])
    P = "Plot_Status!"
    for i, site in enumerate(SITES):
        r = 5 + i
        ss[f"A{r}"] = f"=Settings!E{4 + i}"
        ss[f"B{r}"] = f'=COUNTIF(Planting_Log!$B$5:$B$60,A{r})'
        ss[f"C{r}"] = f'=SUMIFS(Planting_Log!$G$5:$G$60,Planting_Log!$B$5:$B$60,A{r})'
        ss[f"D{r}"] = f'=SUMIFS({mr("K")},{mr("M")},A{r})'
        ss[f"E{r}"] = f'=SUMIFS({P}$G$5:$G$60,{P}$B$5:$B$60,A{r})'
        ss[f"F{r}"] = f'=IF(C{r}=0,"",E{r}/C{r})'
        ss[f"G{r}"] = f'=COUNTIFS({P}$B$5:$B$60,A{r},{P}$L$5:$L$60,"Replant")'
        ss[f"H{r}"] = f'=SUMIFS({P}$M$5:$M$60,{P}$B$5:$B$60,A{r})'
        ss[f"I{r}"] = f'=COUNTIFS({P}$B$5:$B$60,A{r},{P}$N$5:$N$60,">"&OverdueDays)'
        ss[f"J{r}"] = f'=COUNTIFS({mr("M")},A{r},{mr("Z")},">0")'
    ss["A8"] = "All sites"
    for col in "BCDEGHIJ":
        ss[f"{col}8"] = f"=SUM({col}5:{col}7)"
    ss["F8"] = '=IF(C8=0,"",E8/C8)'
    ss["A9"] = "Rows with an unknown plot ID"
    ss["J9"] = f'=COUNTIFS({mr("B")},"<>",{mr("M")},"")'
    style_range(ss, 5, 9, 1, 10, CALC_FILL)
    for r in range(5, 9):
        ss[f"F{r}"].number_format = "0%"
    for c in range(1, 11):
        ss.cell(row=8, column=c).font = Font(name="Calibri", bold=True, size=10)

    ss["A12"] = "Stocking by census month (clean census rows only)"
    ss["A12"].font = Font(name="Calibri", bold=True, size=11, color=ACCENT)
    head_cells = ["Site", "Month 1", "Month 3", "Month 6"]
    for c, text in enumerate(head_cells, start=1):
        cell = ss.cell(row=13, column=c, value=text)
        cell.font, cell.fill, cell.border = HFONT, HEAD_FILL, BOX
    months = [1, 3, 6]
    for i in range(len(SITES)):
        r = 14 + i
        ss[f"A{r}"] = f"=Settings!E{4 + i}"
        for j, m in enumerate(months):
            col = get_column_letter(2 + j)
            ss[f"{col}{r}"] = (f'=IFERROR(AVERAGEIFS({mr("U")},{mr("M")},$A{r},{mr("P")},{m},'
                               f'{mr("D")},"Census",{mr("Z")},0),"")')
            ss[f"{col}{r}"].number_format = "0%"
    style_range(ss, 14, 16, 1, 4, CALC_FILL)
    for r in range(14, 17):
        for col in "BCD":
            ss[f"{col}{r}"].number_format = "0%"

    ss["F12"] = "Seedlings lost by main cause (cause recorded at each plot's last clean visit)"
    ss["F12"].font = Font(name="Calibri", bold=True, size=11, color=ACCENT)
    for c, text in enumerate(["Main cause", "Plots", "Seedlings lost"], start=6):
        cell = ss.cell(row=13, column=c, value=text)
        cell.font, cell.fill, cell.border = HFONT, HEAD_FILL, BOX
    for i, cause in enumerate(CAUSES):
        r = 14 + i
        ss[f"F{r}"] = f"=Settings!G{4 + i}"
        ss[f"G{r}"] = f'=COUNTIF({P}$R$5:$R$60,F{r})'
        ss[f"H{r}"] = f'=SUMIFS({P}$Q$5:$Q$60,{P}$R$5:$R$60,F{r})'
    style_range(ss, 14, 13 + len(CAUSES), 6, 8, CALC_FILL)
    page(ss)

    # ------------------------------------------------------------------ Dashboard
    db = wb.create_sheet("Dashboard", 1)
    title(db, "Dashboard", "Illustrative data. Figures update from Planting_Log and Monitoring; status uses clean rows only.")
    kpis = [
        ("Seedlings planted (design)", "=Site_Summary!C8", "0"),
        ("Replanted to date", "=Site_Summary!D8", "0"),
        ("Alive at last clean visit", "=Site_Summary!E8", "0"),
        ("Stocking vs design", "=Site_Summary!F8", "0%"),
        ("Plots to replant", "=Site_Summary!G8", "0"),
        ("Seedlings to order", "=Site_Summary!H8", "0"),
        ("Plots with census overdue", "=Site_Summary!I8", "0"),
        ("Flagged rows to correct", f'=COUNTIF({mr("Z")},">0")', "0"),
    ]
    db["A3"] = "Report date"
    db["B3"] = "=ReportDate"
    db["B3"].number_format = DATE_FMT
    db["A3"].font = Font(name="Calibri", bold=True, size=10)
    for i, (label, formula, fmt) in enumerate(kpis):
        col = 1 + 2 * (i % 4)
        row = 5 + 3 * (i // 4)
        lab = db.cell(row=row, column=col, value=label)
        lab.font = Font(name="Calibri", size=9, color="475467")
        val = db.cell(row=row + 1, column=col, value=formula)
        val.font = Font(name="Calibri", bold=True, size=18, color=ACCENT)
        val.number_format = fmt
        val.alignment = Alignment(horizontal="left")
    for c in range(1, 9):
        db.column_dimensions[get_column_letter(c)].width = 15

    def axis_style(ch, pct=True, ymax=None):
        ch.x_axis.delete = False
        ch.y_axis.delete = False
        if pct:
            ch.y_axis.number_format = "0%"
            ch.y_axis.scaling.min = 0
            ch.y_axis.scaling.max = ymax if ymax else 1
        ch.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill="E1E0D9"))
        ch.legend.position = "b"
        if ch.title is not None:
            ch.title.overlay = False

    bar = BarChart()
    bar.type = "col"
    bar.title = "Stocking at the last clean visit, by plot (grey line = replanting trigger)"
    bar.add_data(Reference(ps, min_col=9, min_row=4, max_row=pend), titles_from_data=True)
    bar.set_categories(Reference(ps, min_col=1, min_row=5, max_row=pend))
    bar.series[0].graphicalProperties.solidFill = S1
    bar.series[0].graphicalProperties.line.solidFill = S1
    bar.gapWidth = 80
    line = LineChart()
    line.add_data(Reference(ps, min_col=19, min_row=4, max_row=pend), titles_from_data=True)
    line.series[0].graphicalProperties.line.solidFill = GREY
    line.series[0].graphicalProperties.line.width = 19050
    line.series[0].marker.symbol = "none"
    line.series[0].smooth = False
    axis_style(bar)
    bar.legend = None
    bar += line
    bar.height, bar.width = 7.5, 17
    db.add_chart(bar, "A12")

    trend = LineChart()
    trend.title = "Stocking by census month, by site"
    trend.add_data(Reference(ss, min_col=2, min_row=14, max_col=4, max_row=16), from_rows=True, titles_from_data=False)
    trend.set_categories(Reference(ss, min_col=2, min_row=13, max_col=4, max_row=13))
    for s, col, i in zip(trend.series, (S1, S2, S3), range(3)):
        s.graphicalProperties.line.solidFill = col
        s.graphicalProperties.line.width = 28575
        s.marker.symbol = "circle"
        s.marker.size = 7
        s.marker.graphicalProperties.solidFill = col
        s.marker.graphicalProperties.line.solidFill = col
        s.smooth = False
        s.tx = SeriesLabel(strRef=StrRef(f"'Site_Summary'!$A${14 + i}"))
    axis_style(trend, ymax=1)
    trend.y_axis.scaling.min = 0.5
    trend.legend.position = "r"
    trend.legend.overlay = False
    trend.height, trend.width = 7.5, 15
    db.add_chart(trend, "A28")

    cause = BarChart()
    cause.type = "bar"
    cause.title = "Seedlings lost, by main cause at the last clean visit"
    cause.add_data(Reference(ss, min_col=8, min_row=13, max_row=13 + len(CAUSES)), titles_from_data=True)
    cause.set_categories(Reference(ss, min_col=6, min_row=14, max_row=13 + len(CAUSES)))
    cause.series[0].graphicalProperties.solidFill = S1
    cause.series[0].graphicalProperties.line.solidFill = S1
    cause.gapWidth = 70
    axis_style(cause, pct=False)
    cause.legend = None
    cause.x_axis.scaling.orientation = "maxMin"
    cause.height, cause.width = 7.5, 12
    db.add_chart(cause, "F28")
    page(db)
    db.page_setup.fitToHeight = 1

    # ------------------------------------------------------------------ Field_Day_SOP
    sop = wb.create_sheet("Field_Day_SOP")
    title(sop, "Planting day: standard operating procedure (draft)",
          "A draft built from public restoration guidance and my own field work; to be adapted to CarbonEthics' own SOPs.")
    head(sop, 4, ["No.", "Phase", "Step", "What good looks like", "Owner", "Record kept"], [5, 20, 40, 60, 20, 30])
    for i, (phase, step, std, owner, rec) in enumerate(SOP, start=1):
        r = 4 + i
        for c, v in enumerate([i, phase, step, std, owner, rec], start=1):
            sop.cell(row=r, column=c, value=v)
    style_range(sop, 5, 4 + len(SOP), 1, 6, None, wrap=True)
    page(sop)

    # ------------------------------------------------------------------ Risk_Register
    rr = wb.create_sheet("Risk_Register")
    title(rr, "Risk register for planting and monitoring (draft)",
          "Likelihood and severity 1 (low) to 3 (high). Score = likelihood x severity; 6 or 9 = High, 3 or 4 = Medium, 1 or 2 = Low.")
    head(rr, 4, ["ID", "Risk", "Cause or trigger", "Likelihood", "Severity", "Score", "Rating",
                 "Control before the day", "Control on the day", "Owner", "Early warning"],
         [6, 34, 26, 9, 8, 7, 9, 44, 34, 14, 30])
    for i, row in enumerate(RISKS):
        r = 5 + i
        rid, risk, cause_, l, s, before, day, owner, warn = row
        values = [rid, risk, cause_, l, s, f"=D{r}*E{r}",
                  f'=IF(F{r}>=6,"High",IF(F{r}>=3,"Medium","Low"))', before, day, owner, warn]
        for c, v in enumerate(values, start=1):
            rr.cell(row=r, column=c, value=v)
    style_range(rr, 5, 4 + len(RISKS), 1, 11, None, wrap=True)
    for r in range(5, 5 + len(RISKS)):
        for c in (4, 5):
            rr.cell(row=r, column=c).fill = INPUT_FILL
        for c in (6, 7):
            rr.cell(row=r, column=c).fill = CALC_FILL
    rend = 4 + len(RISKS)
    for text, fill in (("High", BAD), ("Medium", WARN), ("Low", GOOD)):
        rr.conditional_formatting.add(f"G5:G{rend}", FormulaRule(formula=[f'$G5="{text}"'], fill=fill))
    page(rr)

    wb.save(out_path)
    print("wrote", out_path)


if __name__ == "__main__":
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).with_name("planting-monitoring-kit-illustrative.xlsx")
    build(out)
